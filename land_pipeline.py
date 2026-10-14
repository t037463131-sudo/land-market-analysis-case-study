"""Import a course/work land export into SQLite; run aggregate SQL checks.
Usage: python land_pipeline.py SOURCE.xlsx --out work_results
Requires openpyxl. Raw database and row-level exports are private working files.
"""
import argparse
import csv
import json
import re
import sqlite3
from datetime import date
from pathlib import Path
from openpyxl import load_workbook


def number(value):
    if value is None or str(value).strip() == "":
        return None
    return float(str(value).replace(",", "").replace("坪", "").strip())


def roc_date(value):
    if value is None:
        return None
    y, m, d = map(int, str(value).strip().split("/"))
    return date(y + 1911, m, d).isoformat()


def import_workbook(source, connection):
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        if not {"案件列表", "土地"}.issubset(workbook.sheetnames):
            raise ValueError("Expected source sheets: 案件列表 and 土地")
        connection.executescript('''
        PRAGMA foreign_keys = ON;
        CREATE TABLE transactions (
            case_id INTEGER PRIMARY KEY, transaction_date TEXT,
            total_twd REAL, reported_unit_wan_ping REAL, area_ping REAL,
            area_m2 REAL, source_note TEXT
        );
        CREATE TABLE parcel_details (
            detail_id TEXT PRIMARY KEY, case_id INTEGER NOT NULL,
            parcel_identifier TEXT, transferred_area_ping REAL, zoning TEXT,
            source_note TEXT,
            FOREIGN KEY(case_id) REFERENCES transactions(case_id)
        );
        CREATE INDEX parcel_case_index ON parcel_details(case_id);
        ''')
        for row in workbook["案件列表"].iter_rows(min_row=3, values_only=True):
            if not row or row[0] is None:
                continue
            case_id = int(str(row[0]).strip())
            area = number(row[6])
            total = number(row[4])
            connection.execute("INSERT INTO transactions VALUES (?,?,?,?,?,?,?)", (
                case_id, roc_date(row[3]), None if total is None else total * 10000,
                number(row[5]), area, None if area is None else area / 0.3025,
                str(row[18] or "").strip()))
        for row in workbook["土地"].iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            detail_id = str(row[0]).strip()
            match = re.fullmatch(r"(\d+)-(\d+)", detail_id)
            if not match:
                raise ValueError(f"Invalid detail ID: {detail_id}")
            connection.execute("INSERT INTO parcel_details VALUES (?,?,?,?,?,?)", (
                detail_id, int(match.group(1)), str(row[1] or "").strip(),
                number(row[2]), str(row[3] or "").strip(), str(row[4] or "").strip()))
        connection.commit()
    finally:
        workbook.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("--out", type=Path, default=Path("work_results"))
    args = parser.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)
    db_path = args.out / "land_analysis.sqlite"
    if db_path.exists():
        raise FileExistsError("Use a new output directory; existing database will not be overwritten")
    with sqlite3.connect(db_path) as connection:
        import_workbook(args.source, connection)
        sql_path = Path(__file__).resolve().with_name("analysis_queries.sql")
        connection.executescript(sql_path.read_text(encoding="utf-8"))
        summary = {
            "transaction_count": connection.execute("SELECT COUNT(*) FROM transactions").fetchone()[0],
            "parcel_detail_row_count": connection.execute("SELECT COUNT(*) FROM parcel_details").fetchone()[0],
            "distinct_parcel_identifier_count": connection.execute(
                "SELECT COUNT(DISTINCT parcel_identifier) FROM parcel_details").fetchone()[0],
            "orphan_parcel_rows": len(connection.execute("PRAGMA foreign_key_check").fetchall()),
            "transactions_outside_rounding_tolerance": connection.execute(
                "SELECT COUNT(*) FROM area_reconciliation WHERE outside_rounding_tolerance = 1").fetchone()[0],
            "rows_with_nonzero_rounded_area_difference": connection.execute(
                "SELECT COUNT(*) FROM area_reconciliation WHERE ABS(area_difference_ping) > 0.000001").fetchone()[0],
        }
        (args.out / "validation_summary.json").write_text(
            json.dumps(summary, indent=2), encoding="utf-8")
        for view in ["area_reconciliation", "screening_flags"]:
            cursor = connection.execute(f"SELECT * FROM {view} ORDER BY case_id")
            with (args.out / f"{view}.csv").open("w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow([d[0] for d in cursor.description])
                writer.writerows(cursor)
        cursor = connection.execute('''SELECT substr(transaction_date,1,4) AS year,
            COUNT(*) AS transaction_count FROM transactions GROUP BY year ORDER BY year''')
        with (args.out / "annual_record_counts.csv").open("w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow([d[0] for d in cursor.description])
            writer.writerows(cursor)
    print(json.dumps(summary, indent=2))
    print("Local database and row-level CSVs may contain identifying source information. Do not upload them.")


if __name__ == "__main__":
    main()
